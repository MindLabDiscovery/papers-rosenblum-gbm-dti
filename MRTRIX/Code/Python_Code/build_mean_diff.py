import pandas as pd
import numpy as np
from scipy import stats
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
import os

# ── EDIT THESE TWO THINGS WHEN YOU ADD NEW PATIENTS ──────────────────────────

tumor_sides = {
    'PT_09': 'right', 'PT_25': 'left',  'PT_33': 'right', 'PT_34': 'left',
    'PT_48': 'right', 'PT_68': 'right', 'PT_71': 'right', 'PT_72': 'right',
    'PT_90': 'left',  'PT_108': 'left', 'PT_117': 'left', 'PT_119': 'right',
    'PT_127': 'left', 'PT_133': 'left', 'PT_201': 'left', 'PT_203': 'right',
    'PT_256': 'left', 'PT_273': 'right', 'PT_275': 'left', 'PT_301': 'right',
    # Add new patients here, e.g.: 'PT_305': 'left',
}

csv_dir     = '/Users/rosenble/Desktop/GBM/Glioblastoma/Data/PT_TRACT_STATS/Plotted_Stats'
output_path = '/Users/rosenble/Desktop/GBM/Glioblastoma/Data/Tract_Overall_Stats/Mean_Difference_Per_Metric.xlsx'

# ─────────────────────────────────────────────────────────────────────────────

categories = {
    'AF':'Association', 'ATR':'Projection', 'CG':'Association', 'CST':'Projection',
    'FPT':'Projection', 'FX':'Commissural', 'ICP':'Projection', 'IFO':'Association',
    'ILF':'Association', 'MLF':'Association', 'OR':'Projection', 'POPT':'Projection',
    'SCP':'Projection', 'SLF_I':'Association', 'SLF_II':'Association', 'SLF_III':'Association',
    'ST_FO':'Projection', 'ST_OCC':'Projection', 'ST_PAR':'Projection', 'ST_POSTC':'Projection',
    'ST_PREC':'Projection', 'ST_PREF':'Projection', 'ST_PREM':'Projection', 'STR':'Projection',
    'T_OCC':'Projection', 'T_PAR':'Projection', 'T_POSTC':'Projection', 'T_PREC':'Projection',
    'T_PREF':'Projection', 'T_PREM':'Projection', 'UF':'Association'
}

metrics       = ['mean_FA', 'mean_MD', 'mean_AD', 'mean_RD']
metric_labels = ['FA', 'MD', 'AD', 'RD']
patients      = sorted(tumor_sides.keys())

# Read all CSVs
data = {}
for pt in patients:
    f = os.path.join(csv_dir, f'{pt}.csv')
    if os.path.exists(f):
        df = pd.read_csv(f)
        df = df.set_index('tract')
        data[pt] = df
    else:
        print(f'WARNING: {pt}.csv not found in {csv_dir}')

# Get bilateral tracts
sample = list(data.values())[0]
left_tracts = [t.replace('_left', '') for t in sample.index if t.endswith('_left')]
bilateral   = [t for t in left_tracts if f'{t}_right' in sample.index and t in categories]
print(f'Bilateral tracts found: {len(bilateral)}')

with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
    for metric, label in zip(metrics, metric_labels):
        ipsi_cols   = [f'{pt}_Ipsi'  for pt in patients]
        contra_cols = [f'{pt}_Contra' for pt in patients]
        rows = []
        for tract in bilateral:
            ipsi_vals, contra_vals = [], []
            for pt in patients:
                if pt not in data:
                    ipsi_vals.append(np.nan)
                    contra_vals.append(np.nan)
                    continue
                side        = tumor_sides[pt]
                ipsi_name   = f'{tract}_{side}'
                contra_name = f'{tract}_{"left" if side == "right" else "right"}'
                try:
                    ipsi_vals.append(data[pt].loc[ipsi_name, metric])
                    contra_vals.append(data[pt].loc[contra_name, metric])
                except:
                    ipsi_vals.append(np.nan)
                    contra_vals.append(np.nan)

            diffs = [c - i for c, i in zip(contra_vals, ipsi_vals)
                     if not (np.isnan(c) or np.isnan(i))]
            if len(diffs) >= 3:
                _, pval = stats.wilcoxon(diffs)
            else:
                pval = np.nan
            mean_diff = np.nanmean(diffs)
            sem       = np.nanstd(diffs, ddof=1) / np.sqrt(len(diffs))

            row = {'Tract': tract, 'Category': categories[tract],
                   'Mean_Diff': mean_diff, 'SEM': sem, 'p_value': pval}
            for pt, iv, cv in zip(patients, ipsi_vals, contra_vals):
                row[f'{pt}_Ipsi']  = iv
                row[f'{pt}_Contra'] = cv
            rows.append(row)

        df_out = pd.DataFrame(rows)

        # BH FDR correction
        pvals = df_out['p_value'].values
        n     = len(pvals)
        order = np.argsort(pvals)
        adj   = pvals[order] * n / (np.arange(n) + 1)
        for k in range(n - 2, -1, -1):
            adj[k] = min(adj[k], adj[k + 1])
        fdr = np.empty(n)
        fdr[order] = np.minimum(adj, 1)
        df_out['p_FDR'] = fdr

        # Sort descending by Mean_Diff
        df_out = df_out.sort_values('Mean_Diff', ascending=False).reset_index(drop=True)

        # Column order
        fixed = ['Tract', 'Category', 'Mean_Diff', 'SEM', 'p_value', 'p_FDR']
        cols  = fixed + [c for c in ipsi_cols + contra_cols if c in df_out.columns]
        df_out[cols].to_excel(writer, sheet_name=label, index=False)

# Green highlight for p<0.05
wb    = load_workbook(output_path)
green = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
for label in metric_labels:
    ws      = wb[label]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    p_col   = headers.index('p_value') + 1
    for row in range(2, ws.max_row + 1):
        val = ws.cell(row, p_col).value
        if val is not None and val < 0.05:
            for c in range(1, ws.max_column + 1):
                ws.cell(row, c).fill = green
wb.save(output_path)

print(f'\nSaved: {output_path}')
print(f'Patients included: {len(data)} of {len(patients)}')
for label in metric_labels:
    df  = pd.read_excel(output_path, sheet_name=label)
    sig = (df['p_value'] < 0.05).sum()
    fdr = (df['p_FDR'] < 0.05).sum()
    print(f'  {label}: {sig} significant (p<0.05), {fdr} after FDR')
