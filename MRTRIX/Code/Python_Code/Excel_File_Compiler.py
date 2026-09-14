"""
GBM DTI Excel Populator
========================
Scans a folder of patient CSV files (e.g. PT_09.csv, PT_25.csv),
extracts FA/MD/AD/RD tract values, and writes them into GBM_DTI_Master.xlsx.

Asymmetry, composite, and tract count columns calculate automatically in Excel.
Tumor Side, Survival, and Status still need to be entered manually in Master sheet.

REQUIREMENTS
------------
pip install openpyxl

HOW TO USE
----------
1. Edit FOLDER and EXCEL_FILE below to match your paths
2. Run:  python populate_gbm_excel.py
"""

import os
import csv
import re
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl not found. Run:  pip install openpyxl")

# ── EDIT THESE TWO PATHS ──────────────────────────────────────────────────────
FOLDER     = r"/Users/rosenble/Desktop/GBM/Glioblastoma/Data/PT_TRACT_STATS/Plotted_Stats"         # folder containing PT_XX.csv files
EXCEL_FILE = r"/Users/rosenble/Desktop/GBM/Glioblastoma/Data/Tract_Overall_Stats_V2/GBM_DTI_Master_Compiled.xlsx"     # your master spreadsheet
# ─────────────────────────────────────────────────────────────────────────────

TRACTS = [
    'AF','ATR','CG','CST','FPT','FX','ICP','IFO','ILF','MLF',
    'OR','POPT','SCP','SLF_I','SLF_II','SLF_III','ST_FO','ST_OCC',
    'ST_PAR','ST_POSTC','ST_PREC','ST_PREF','ST_PREM','STR','T_OCC',
    'T_PAR','T_POSTC','T_PREC','T_PREF','T_PREM','UF'
]
METRICS = ['FA','MD','AD','RD']
CSV_COL = {'FA':'mean_FA','MD':'mean_MD','AD':'mean_AD','RD':'mean_RD'}
DATA_START_ROW = 2

def patient_sort_key(filename):
    match = re.search(r'(\d+)', filename)
    return int(match.group(1)) if match else float('inf')

def read_csv(filepath):
    data = {}
    with open(filepath, newline='') as f:
        reader = csv.DictReader(f)
        for row in reader:
            tract = row.get('tract','').strip()
            if not tract:
                continue
            values = {}
            ok = True
            for metric, col in CSV_COL.items():
                try:
                    values[metric] = float(row[col])
                except (KeyError, ValueError):
                    ok = False
                    break
            if ok:
                data[tract] = values
    return data

def build_tract_lookup(csv_data):
    lookup = {}
    for tract in TRACTS:
        left_key  = f'{tract}_left'
        right_key = f'{tract}_right'
        left_data = right_data = None
        for k, v in csv_data.items():
            if k.lower() == left_key.lower():
                left_data = v
            elif k.lower() == right_key.lower():
                right_data = v
        if left_data or right_data:
            lookup[tract] = {'left': left_data, 'right': right_data}
    return lookup

def populate_excel(csv_folder, excel_path):
    all_files = [f for f in os.listdir(csv_folder)
                 if f.lower().endswith('.csv')
                 and re.search(r'pt.?\d+', f, re.IGNORECASE)]
    if not all_files:
        sys.exit(f"No patient CSV files found in: {csv_folder}")
    all_files.sort(key=patient_sort_key)
    print(f"Found {len(all_files)} patient CSV files")

    wb = openpyxl.load_workbook(excel_path)
    ws_master = wb['Master']
    metric_sheets = {m: wb[m] for m in METRICS}
    max_row = DATA_START_ROW + 49

    for r in range(DATA_START_ROW, max_row + 1):
        ws_master.cell(r, 1).value = None
    for metric, ws in metric_sheets.items():
        for ti in range(len(TRACTS)):
            for r in range(DATA_START_ROW, max_row + 1):
                ws.cell(r, 3 + ti*3).value = None
                ws.cell(r, 4 + ti*3).value = None

    loaded = 0
    skipped = []
    for i, filename in enumerate(all_files):
        row = DATA_START_ROW + i
        if row > max_row:
            print(f"WARNING: More than 50 patients — {filename} and beyond skipped.")
            break
        filepath = os.path.join(csv_folder, filename)
        patient_id = re.sub(r'\.csv$', '', filename, flags=re.IGNORECASE)
        try:
            csv_data = read_csv(filepath)
        except Exception as e:
            skipped.append((filename, str(e))); continue
        if not csv_data:
            skipped.append((filename, "no readable data")); continue
        tract_lookup = build_tract_lookup(csv_data)
        ws_master.cell(row, 1).value = patient_id
        for metric, ws in metric_sheets.items():
            for ti, tract in enumerate(TRACTS):
                if tract in tract_lookup:
                    left_data  = tract_lookup[tract].get('left')
                    right_data = tract_lookup[tract].get('right')
                    if left_data:
                        ws.cell(row, 3+ti*3).value = left_data[metric]
                    if right_data:
                        ws.cell(row, 4+ti*3).value = right_data[metric]
        loaded += 1
        print(f"  ✓ {patient_id} → row {row} ({len(tract_lookup)}/{len(TRACTS)} tracts)")

    wb.save(excel_path)
    print(f"\nDone — {loaded} patients loaded")
    if skipped:
        print(f"Skipped: {[f for f,_ in skipped]}")

if __name__ == '__main__':
    if len(sys.argv) == 3:
        FOLDER, EXCEL_FILE = sys.argv[1], sys.argv[2]
    populate_excel(FOLDER, EXCEL_FILE)